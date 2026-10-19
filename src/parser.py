import os
from typing import Optional
from docx import Document
from PyPDF2 import PdfReader
import io


class ResumeParser:
    """Parse resumes from PDF and Word documents."""
    
    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """Extract text from PDF file."""
        try:
            with open(file_path, 'rb') as file:
                reader = PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text()
                return text
        except Exception as e:
            raise Exception(f"Error parsing PDF: {str(e)}")
    
    @staticmethod
    def parse_docx(file_path: str) -> str:
        """Extract text from Word document."""
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
    
    @staticmethod
    def parse_file(file_path: str) -> str:
        """Parse resume file based on extension."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            return ResumeParser.parse_pdf(file_path)
        elif file_extension in ['.docx', '.doc']:
            return ResumeParser.parse_docx(file_path)
        elif file_extension == '.txt':
            for enc in ['utf-8', 'utf-8-sig', 'cp1252', 'latin-1']:
                try:
                    with open(file_path, 'r', encoding=enc) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
    
    @staticmethod
    def parse_uploaded_file(file_content: bytes, filename: str) -> str:
        """Parse uploaded file content."""
        file_extension = os.path.splitext(filename)[1].lower()
        
        if file_extension == '.pdf':
            file_like = io.BytesIO(file_content)
            reader = PdfReader(file_like)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
            return text
        elif file_extension in ['.docx', '.doc']:
            file_like = io.BytesIO(file_content)
            doc = Document(file_like)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        elif file_extension == '.txt':
            for enc in ['utf-8', 'utf-8-sig', 'cp1252', 'latin-1']:
                try:
                    return file_content.decode(enc)
                except UnicodeDecodeError:
                    continue
            return file_content.decode('utf-8', errors='replace')
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
